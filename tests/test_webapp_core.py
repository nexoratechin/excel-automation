from __future__ import annotations

import io
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.config import ConfigError
from src.webapp_core import MAX_BYTES, inspect_file, process_file

ROWS = [
    ("Rahul Sharma", "rahul.sharma@gmail.com", "9876543210"),
    ("Priya Patel", "Priya.Patel@Yahoo.com", "+91 98765 43211"),
    ("Rahul Again", "RAHUL.SHARMA@GMAIL.COM", "9876543210"),
    ("Karan Mehta", "karan.mehta@gmail", "9876543212"),
    ("Sameer Khan", "", "9876543213"),
    ("Anjali Desai", "anjali@gmail.com", ""),
    (None, None, None),
]
COLUMNS = ["Name", "Email Address", "Mobile No"]


def dirty_csv() -> bytes:
    df = pd.DataFrame(ROWS, columns=COLUMNS)
    return df.to_csv(index=False).encode("utf-8")


def dirty_xlsx() -> bytes:
    df = pd.DataFrame(ROWS, columns=COLUMNS)
    buffer = io.BytesIO()
    df.to_excel(buffer, index=False)
    return buffer.getvalue()


class TestInspect:
    def test_detects_columns_from_content(self):
        info = inspect_file(dirty_csv(), "list.csv")
        assert info.columns == COLUMNS
        assert info.email_guess == "Email Address"
        assert info.phone_guess == "Mobile No"

    def test_counts_rows(self):
        assert inspect_file(dirty_csv(), "list.csv").row_count == 7

    def test_preview_is_small(self):
        assert len(inspect_file(dirty_csv(), "list.csv").preview) == 5


class TestProcess:
    def test_returns_a_valid_workbook(self, tmp_path):
        result = process_file(dirty_csv(), "contacts.csv")
        path = tmp_path / "out.xlsx"
        path.write_bytes(result.workbook)

        assert path.read_bytes()[:2] == b"PK"  # xlsx is a zip
        assert load_workbook(path).sheetnames == ["Report", "Cleaned", "Rejected", "Changes"]

    def test_counts_are_correct(self):
        result = process_file(dirty_csv(), "contacts.csv")
        assert result.cleaned_total == 4
        assert result.rejected_total == 2   # 1 duplicate + 1 bad email
        assert result.stats.duplicates_removed == 1
        assert result.stats.invalid_email == 1
        assert result.stats.blank_rows == 1

    def test_xlsx_upload_works(self):
        assert len(process_file(dirty_xlsx(), "contacts.xlsx").workbook) > 1000

    def test_report_shows_uploaded_filename(self, tmp_path):
        result = process_file(dirty_csv(), "client_list.csv")
        path = tmp_path / "out.xlsx"
        path.write_bytes(result.workbook)
        text = " ".join(
            str(c.value)
            for row in load_workbook(path)["Report"].iter_rows()
            for c in row
            if c.value is not None
        )
        assert "client_list.csv" in text
        assert "contacts.csv" not in text

    def test_explicit_columns_win(self):
        df = pd.DataFrame([("a@gmail.com", "9876543210")], columns=["A", "B"])
        result = process_file(
            df.to_csv(index=False).encode(),
            "x.csv",
            email_col="A",
            phone_col="B",
        )
        assert result.email_col == "A"
        assert result.phone_col == "B"
        assert result.cleaned_total == 1

    def test_required_column_rejects_blanks(self):
        result = process_file(dirty_csv(), "contacts.csv", required=["Name"])
        # Sameer has no email but a name, so he survives either way.
        assert result.cleaned_total == 4

    def test_keep_duplicates_reports_without_removing(self):
        result = process_file(dirty_csv(), "contacts.csv", keep_duplicates=True)
        assert result.cleaned_total == 5
        assert result.stats.duplicates_removed == 1

    def test_any_country_relaxes_phone_rule(self):
        result = process_file(dirty_csv(), "contacts.csv", country="ANY")
        assert result.stats.invalid_phone == 0

    def test_phone_format_e164(self):
        result = process_file(dirty_csv(), "contacts.csv", phone_format="e164")
        assert "+919876543210" in result.cleaned_preview["Mobile No"].tolist()


class TestGuards:
    def test_rejects_wrong_file_type(self):
        with pytest.raises(ConfigError, match="not a supported file"):
            process_file(b"hello", "notes.pdf")

    def test_rejects_empty_file(self):
        with pytest.raises(ConfigError, match="empty"):
            process_file(b"", "list.csv")

    def test_rejects_oversized_file(self, monkeypatch):
        import src.webapp_core as core

        monkeypatch.setattr(core, "MAX_BYTES", 4)
        with pytest.raises(ConfigError, match="15 MB|larger"):
            process_file(b"a,b\n1,2\n", "list.csv")

    def test_rejects_file_with_no_rows(self):
        with pytest.raises(ConfigError, match="No data rows"):
            process_file(b"Name,Email\n", "list.csv")

    def test_max_bytes_is_fifteen_mb(self):
        assert MAX_BYTES == 15 * 1024 * 1024
