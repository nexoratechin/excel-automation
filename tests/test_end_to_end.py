from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

DIRTY = pd.DataFrame(
    [
        ("Rahul Sharma", "rahul.sharma@gmail.com", "9876543210"),
        ("Priya Patel", "Priya.Patel@Yahoo.com", "+91 98765 43299"),
        ("Rahul Again", "RAHUL.SHARMA@GMAIL.COM", "9876543210"),
        ("Neha Gupta", "neha@gmail.com", "9876543211"),
        ("Neha Dup", "neha2@gmail.com", "09876543211"),
        ("Karan Mehta", "karan.mehta@gmail", "9876543212"),
        ("Meera Joshi", "meera@gmail.com", "12345"),
        ("Sameer Khan", "", "9876543213"),
        ("Anjali Desai", "anjali@gmail.com", ""),
        (None, None, None),
    ],
    columns=["Name", "Email Address", "Mobile No"],
)


@pytest.fixture
def dirty_xlsx(tmp_path):
    path = tmp_path / "dirty.xlsx"
    DIRTY.to_excel(path, index=False, sheet_name="Contacts")
    return path


def run(args, capsys):
    from src.main import main

    code = main(args)
    out = capsys.readouterr()
    return code, out.out + out.err


def test_end_to_end_produces_three_sheets(dirty_xlsx, tmp_path, capsys):
    out = tmp_path / "clean.xlsx"
    code, _ = run(
        ["--input", str(dirty_xlsx), "--out", str(out),
         "--email-col", "Email Address", "--phone-col", "Mobile No",
         "--required", "Name"],
        capsys,
    )
    assert code == 0
    assert out.exists()

    book = load_workbook(out)
    assert book.sheetnames == ["Report", "Cleaned", "Rejected"]
    assert book["Cleaned"].max_row == 1 + 5     # 5 survivors
    assert book["Rejected"].max_row == 1 + 4    # 2 duplicates + 1 bad email + 1 bad phone


def test_input_file_is_never_modified(dirty_xlsx, tmp_path, capsys):
    before = dirty_xlsx.read_bytes()
    out = tmp_path / "clean.xlsx"
    code, _ = run(["--input", str(dirty_xlsx), "--out", str(out),
                   "--email-col", "Email Address", "--phone-col", "Mobile No"], capsys)
    assert code == 0
    assert dirty_xlsx.read_bytes() == before


def test_refuses_to_overwrite_input(dirty_xlsx, capsys):
    code, output = run(["--input", str(dirty_xlsx), "--out", str(dirty_xlsx),
                        "--email-col", "Email Address"], capsys)
    assert code == 2
    assert "never modifies" in output


def test_reports_duplicates_in_summary(dirty_xlsx, tmp_path, capsys):
    out = tmp_path / "clean.xlsx"
    _, output = run(["--input", str(dirty_xlsx), "--out", str(out),
                     "--email-col", "Email Address", "--phone-col", "Mobile No"], capsys)
    assert "rows kept" in output
    assert "duplicates" in output


def test_csv_input_works(tmp_path, capsys):
    src = tmp_path / "list.csv"
    DIRTY.to_csv(src, index=False)
    out = tmp_path / "clean.xlsx"
    code, _ = run(["--input", str(src), "--out", str(out),
                   "--email-col", "Email Address", "--phone-col", "Mobile No"], capsys)
    assert code == 0
    assert load_workbook(out)["Cleaned"].max_row == 6

def test_config_file_supplies_mapping(dirty_xlsx, tmp_path, capsys):
    cfg = tmp_path / "cfg.json"
    cfg.write_text(
        '{"email_col": "Email Address", "phone_col": "Mobile No", "required": ["Name"]}',
        encoding="utf-8",
    )
    out = tmp_path / "clean.xlsx"
    code, _ = run(["--input", str(dirty_xlsx), "--out", str(out), "--config", str(cfg)], capsys)
    assert code == 0


def test_cli_overrides_config_file(dirty_xlsx, tmp_path, capsys):
    cfg = tmp_path / "cfg.json"
    cfg.write_text('{"email_col": "Wrong Column"}', encoding="utf-8")
    out = tmp_path / "clean.xlsx"
    code, _ = run(
        ["--input", str(dirty_xlsx), "--out", str(out),
         "--config", str(cfg), "--email-col", "Email Address"],
        capsys,
    )
    assert code == 0


def test_missing_column_gives_actionable_error(dirty_xlsx, tmp_path, capsys):
    code, output = run(["--input", str(dirty_xlsx), "--out", str(tmp_path / "c.xlsx"),
                        "--email-col", "Wrong"], capsys)
    assert code == 2
    assert "Columns in this file" in output


def test_list_sheets(dirty_xlsx, capsys):
    code, output = run(["--input", str(dirty_xlsx), "--list-sheets"], capsys)
    assert code == 0
    assert "Contacts" in output


def test_missing_input_is_reported(capsys):
    code, output = run(["--out", "x.xlsx"], capsys)
    assert code == 2
    assert "No input file" in output


def test_report_sheet_contains_counts(dirty_xlsx, tmp_path, capsys):
    out = tmp_path / "clean.xlsx"
    run(["--input", str(dirty_xlsx), "--out", str(out),
         "--email-col", "Email Address", "--phone-col", "Mobile No"], capsys)
    ws = load_workbook(out)["Report"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "Data Cleanup Report" in text
    assert "Rows rejected" in text
    assert "duplicates" in text.lower()
