"""Write the deliverable workbook: a styled, self-explanatory Excel file."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .cleaner import CleanResult
from .config import Config

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=16, color="1F3864")
SECTION_FONT = Font(bold=True, size=12, color="1F3864")
LABEL_FILL = PatternFill("solid", fgColor="F2F2F2")
GOOD_FONT = Font(color="1E7B34", bold=True)
BAD_FONT = Font(color="C00000", bold=True)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MAX_WIDTH = 55


def write_workbook(result: CleanResult, cfg: Config, out_path: str | Path) -> Path:
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        result.cleaned.to_excel(writer, sheet_name="Cleaned", index=False)
        result.rejected.to_excel(writer, sheet_name="Rejected", index=False)

        report = writer.book.create_sheet("Report", 0)
        _fill_report(report, result, cfg)

        _style_data_sheet(writer.sheets["Cleaned"])
        _style_data_sheet(writer.sheets["Rejected"], reason_column=2)
        _style_report(report)
        writer.book.active = 0

    return out


def _fill_report(ws, result: CleanResult, cfg: Config) -> None:
    stats = result.stats
    source = Path(cfg.input_path).name
    _write(ws, 1, 1, "Data Cleanup Report", TITLE_FONT)
    _write(ws, 2, 1, "Source file")
    _write(ws, 2, 2, source)
    _write(ws, 3, 1, "Generated")
    _write(ws, 3, 2, datetime.now().strftime("%d %b %Y, %H:%M"))
    _write(ws, 4, 1, "Rules applied")
    rule_bits = [f"country {cfg.country}", f"phone format {cfg.phone_format}"]
    if cfg.email_col:
        rule_bits.append(f"email column '{cfg.email_col}'")
    if cfg.phone_col:
        rule_bits.append(f"phone column '{cfg.phone_col}'")
    if cfg.required:
        rule_bits.append(f"required: {', '.join(cfg.required)}")
    _write(ws, 4, 2, "; ".join(rule_bits))

    row = 6
    _write(ws, row, 1, "Summary", SECTION_FONT)
    row += 1
    summary_rows = [
        ("Rows read", stats.total_rows, None),
        ("Rows kept", stats.kept, GOOD_FONT),
        ("Rows rejected", stats.rejected, BAD_FONT if stats.rejected else GOOD_FONT),
        ("Duplicates removed", stats.duplicates_removed, BAD_FONT if stats.duplicates_removed else None),
        ("Blank rows dropped", stats.blank_rows, None),
        ("Invalid emails", stats.invalid_email, BAD_FONT if stats.invalid_email else None),
        ("Invalid phone numbers", stats.invalid_phone, BAD_FONT if stats.invalid_phone else None),
        ("Missing required fields", stats.missing_required, BAD_FONT if stats.missing_required else None),
        ("Blank emails (allowed)", stats.blank_email, None),
        ("Blank phones (allowed)", stats.blank_phone, None),
    ]
    for label, value, font in summary_rows:
        _write(ws, row, 1, label, font=None, fill=LABEL_FILL, border=True)
        cell = _write(ws, row, 2, value, font=font, border=True)
        cell.alignment = Alignment(horizontal="right")
        row += 1

    if stats.reason_counts:
        row += 1
        _write(ws, row, 1, "Why rows were rejected", SECTION_FONT)
        row += 1
        for reason, count in sorted(stats.reason_counts.items(), key=lambda kv: -kv[1]):
            _write(ws, row, 1, reason, fill=LABEL_FILL, border=True)
            _write(ws, row, 2, count, border=True).alignment = Alignment(horizontal="right")
            row += 1

    row += 1
    _write(ws, row, 1, "How to read this file", SECTION_FONT)
    row += 1
    notes = [
        "Cleaned - rows that passed every check, with emails and phones normalised.",
        "Rejected - nothing was deleted silently; every dropped row is listed with a reason.",
        "Your original file has not been modified.",
    ]
    for note in notes:
        _write(ws, row, 1, f"- {note}")
        row += 1


def _write(ws, row: int, col: int, value, font=None, fill=None, border: bool = False):
    cell = ws.cell(row=row, column=col, value=value)
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if border:
        cell.border = BORDER
    return cell


def _style_data_sheet(ws, reason_column: int | None = None) -> None:
    if ws.max_row < 1:
        return
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions

    widths: dict[int, int] = {}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 500)):
        for cell in row:
            text = "" if cell.value is None else str(cell.value)
            widths[cell.column] = max(widths.get(cell.column, 0), min(len(text), MAX_WIDTH))
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = max(10, width + 3)

    if reason_column:
        for row in ws.iter_rows(min_row=2, max_col=reason_column):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=cell.column == reason_column)
        ws.column_dimensions[get_column_letter(reason_column)].width = 48


def _style_report(ws) -> None:
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 70
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=2):
        for cell in row:
            if cell.column == 2 and isinstance(cell.value, str) and cell.font.size == 11:
                cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.sheet_view.showGridLines = False
